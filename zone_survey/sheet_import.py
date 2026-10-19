"""
'2025 구역편성(0905)' 시트를 읽어 구역편성 설문 DB(zone_survey.db)를 만든다.

사용법:
    python -m zone_survey.sheet_import <xlsx 경로 또는 구글시트 URL> [--sheet 시트이름] [--reset]

이미 구역장들이 입력한 데이터가 있으면 --reset 없이는 덮어쓰지 않는다.
"""
import argparse
import json
import os
import re
import secrets
import sqlite3
import sys
import tempfile
import urllib.request

import openpyxl

from zone_survey.db import DATABASE, init_schema

DEFAULT_SHEET = '2025 구역편성(0905)'

PHONE_RE = re.compile(r'^0\d{1,2}-?\d{3,4}-?\d{4}$')
PERSON_RE = re.compile(r'^\s*([가-힣A-Za-z]+[A-Z]?)\s*(?:\((\d{2})\))?\s*$')
CHILD_RE = re.compile(r'([가-힣A-Za-z]+)\s*\(([^)]*)\)')
GRADE_RE = re.compile(r'^\s*([A-D])\s*$')


def _s(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def is_phone(v):
    return bool(PHONE_RE.match(_s(v)))


def looks_like_children(v):
    s = _s(v)
    return bool(re.search(r'\([^)]*(남|여)', s)) or '자녀' in s


def parse_person(v):
    """'김민웅(86)' -> ('김민웅', '1986')"""
    s = _s(v)
    m = PERSON_RE.match(s)
    if not m:
        return s, ''
    name, yy = m.group(1), m.group(2)
    return name, expand_year(yy) if yy else ''


def expand_year(yy):
    yy = int(yy)
    return str(1900 + yy if yy > 30 else 2000 + yy)


def parse_children(text):
    """'김다윤(17,남) 김하윤(19,여)' -> [{'name','year','gender','note'}]"""
    s = _s(text)
    if not s or '없음' in s.replace(' ', '') or s in ('0',):
        return []
    children = []
    for name, inner in CHILD_RE.findall(s):
        year, gender, notes = '', '', []
        for tok in re.split(r'[,./\s]+', inner):
            if not tok:
                continue
            if re.fullmatch(r'\d{2}', tok):
                year = tok
            elif tok in ('남', '여'):
                gender = tok
            else:
                notes.append(tok)
        children.append({'name': name, 'year': year, 'gender': gender, 'note': ' '.join(notes)})
    if not children:
        children.append({'name': s, 'year': '', 'gender': '', 'note': ''})
    return children


def parse_attendance(v):
    """시트의 출석 값 -> ('84%' | 'A' | '', 메모로 보낼 문자열)"""
    if v is None:
        return '', ''
    if isinstance(v, (int, float)):
        return f'{round(v * 100)}%', ''
    s = _s(v)
    m = GRADE_RE.match(s)
    if m:
        return m.group(1), ''
    if s in ('', '-', '#N/A'):
        return '', ''
    return '', s


POSITIONAL = ['이름', '전화번호', '이름', '전화번호', '자녀']


def read_blocks(ws):
    """시트를 가정 블록 단위(행 범위)로 나눈다. 모든 구역 열이 같은 행 구조를 공유한다."""
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(rows) if _s(r[2]) == '1구역')
    header = rows[header_idx]
    zone_cols = [(c, _s(header[c])) for c in range(2, len(header)) if _s(header[c])]

    blocks = []  # (role, [(label, row)])
    role = ''
    cur = None
    for r in rows[header_idx + 1:]:
        a, label = _s(r[0]), _s(r[1])
        cells = [_s(r[c]) for c, _ in zone_cols]
        if a:
            role = a.strip()
        love_marker = label == '사랑구역' or (not label and '사랑구역' in cells)
        if love_marker:
            role = '사랑구역'
            cur = None
            continue
        if not label and not any(cells):
            continue
        pos = len(cur[1]) if cur else 0
        if not label:
            label = POSITIONAL[pos] if pos < len(POSITIONAL) else '기타'
        if label == '이름' and (cur is None or pos >= 3):
            cur = (role, [])
            blocks.append(cur)
        if cur is None:
            continue
        cur[1].append((label.replace(' ', ''), r))
    return zone_cols, blocks


def build_family(role, block, col):
    slots = {'names': [], 'phones': [], 'children': '', 'att_m': [], 'att_f': [], 'memo': [], 'other': []}
    name_i = phone_i = 0
    raw = [None, None, None, None]  # husband name, husband phone, wife name, wife phone
    for label, r in block:
        v = r[col] if col < len(r) else None
        s = _s(v)
        if label == '이름':
            if name_i < 2:
                raw[name_i * 2] = s
            name_i += 1
        elif label == '전화번호':
            if phone_i < 2:
                raw[phone_i * 2 + 1] = s
            phone_i += 1
        elif label == '자녀':
            if s and s != '0':
                slots['children'] = s
        elif label in ('남출석', '여출석', '기타'):
            att, memo = parse_attendance(v)
            if label == '남출석' or (label == '기타' and att and not slots['att_m']):
                if att:
                    slots['att_m'].append(att)
            elif att:
                slots['att_f'].append(att)
            if memo:
                slots['memo'].append(memo)
        elif label in ('메모', '비고'):
            if s:
                slots['memo'].append(s)

    # 셀이 한 칸씩 밀려 입력된 경우 바로잡기
    persons = []
    leftover_children = []
    for name_v, phone_v in ((raw[0], raw[1]), (raw[2], raw[3])):
        name_v, phone_v = name_v or '', phone_v or ''
        if looks_like_children(name_v):
            leftover_children.append(name_v)
            name_v = ''
        if looks_like_children(phone_v):
            leftover_children.append(phone_v)
            phone_v = ''
        if is_phone(name_v) and not phone_v:
            name_v, phone_v = '', name_v
        if not name_v and phone_v and not is_phone(phone_v) and PERSON_RE.match(phone_v):
            name_v, phone_v = phone_v, ''
        persons.append([name_v, phone_v])
    # 이름 없이 번호만 있는 사람 -> 앞사람 번호로
    if not persons[1][0] and persons[1][1] and persons[0][0] and not persons[0][1]:
        persons[0][1], persons[1][1] = persons[1][1], ''
    if not persons[1][0] and is_phone(persons[1][1]) and persons[0][0] and not is_phone(persons[0][1]):
        persons[0][1], persons[1][1] = persons[1][1], ''

    children_text = slots['children'] or ' '.join(leftover_children)
    if slots['children'] and leftover_children:
        children_text = slots['children'] + ' ' + ' '.join(leftover_children)

    out = {'role': role.strip() or '구역원', 'children_raw': children_text,
           'children': parse_children(children_text),
           'ref_att_m': slots['att_m'][0] if slots['att_m'] else '',
           'ref_att_f': slots['att_f'][0] if slots['att_f'] else '',
           'sheet_memo': '\n'.join(slots['memo'])}
    for key, (name_v, phone_v) in zip(('husband', 'wife'), persons):
        name, birth = parse_person(name_v)
        out[f'{key}_name'] = name
        out[f'{key}_birth'] = birth
        out[f'{key}_phone'] = phone_v if (not phone_v or is_phone(phone_v)) else ''
        if phone_v and not is_phone(phone_v):
            out['sheet_memo'] = (out['sheet_memo'] + f'\n연락처 칸: {phone_v}').strip()
    if not (out['husband_name'] or out['wife_name']):
        return None
    return out


def load_workbook(source):
    if source.startswith('http'):
        m = re.search(r'/d/([A-Za-z0-9_-]+)', source)
        sheet_id = m.group(1) if m else source
        url = f'https://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx'
        tmp = tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False)
        with urllib.request.urlopen(url) as resp:
            tmp.write(resp.read())
        tmp.close()
        source = tmp.name
    return openpyxl.load_workbook(source, data_only=True)


def parse_sheet(ws):
    zone_cols, blocks = read_blocks(ws)
    zones = []
    for order, (col, zone_label) in enumerate(zone_cols):
        families = []
        leaders = []
        for role, block in blocks:
            fam = build_family(role, block, col)
            if not fam:
                continue
            if zone_label == '임원':
                fam['role'] = '임원'
            families.append(fam)
            if fam['role'] == '구역장' or (zone_label == '임원' and not leaders):
                leaders = [n for n in (fam['husband_name'], fam['wife_name']) if n]
        code = zone_label.replace('구역', '') or zone_label
        zones.append({'code': code, 'name': zone_label, 'sort': order,
                      'leader_names': ' · '.join(leaders), 'families': families})
    return zones


def import_zones(zones, reset=False):
    conn = sqlite3.connect(DATABASE)
    init_schema(conn)
    existing = conn.execute('SELECT COUNT(*) FROM zone_families WHERE updated_at IS NOT NULL').fetchone()[0]
    if existing and not reset:
        conn.close()
        sys.exit(f'이미 구역장이 작성한 가정이 {existing}건 있습니다. 덮어쓰려면 --reset 을 붙이세요.')
    conn.execute('DELETE FROM zone_families')
    conn.execute('DELETE FROM zones')
    for z in zones:
        cur = conn.execute(
            'INSERT INTO zones (code, name, sort, leader_names, token) VALUES (?, ?, ?, ?, ?)',
            (z['code'], z['name'], z['sort'], z['leader_names'], secrets.token_urlsafe(8)))
        zone_id = cur.lastrowid
        for i, f in enumerate(z['families']):
            original = {k: f[k] for k in ('husband_name', 'husband_birth', 'husband_phone',
                                          'wife_name', 'wife_birth', 'wife_phone', 'children')}
            conn.execute(
                '''INSERT INTO zone_families
                   (zone_id, sort, role_prev, husband_name, husband_birth, husband_phone,
                    wife_name, wife_birth, wife_phone, children_json, children_raw,
                    ref_att_m, ref_att_f, sheet_memo, original_json)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                (zone_id, i, f['role'], f['husband_name'], f['husband_birth'], f['husband_phone'],
                 f['wife_name'], f['wife_birth'], f['wife_phone'],
                 json.dumps(f['children'], ensure_ascii=False), f['children_raw'],
                 f['ref_att_m'], f['ref_att_f'], f['sheet_memo'],
                 json.dumps(original, ensure_ascii=False)))
    conn.commit()
    conn.close()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('source', help='xlsx 파일 경로 또는 구글시트 URL')
    ap.add_argument('--sheet', default=DEFAULT_SHEET)
    ap.add_argument('--reset', action='store_true', help='기존 입력 데이터를 지우고 새로 불러오기')
    ap.add_argument('--dry-run', action='store_true', help='DB에 쓰지 않고 파싱 결과만 출력')
    args = ap.parse_args()

    wb = load_workbook(args.source)
    zones = parse_sheet(wb[args.sheet])
    for z in zones:
        print(f"[{z['name']}] {z['leader_names']} - {len(z['families'])}가정")
        if args.dry_run:
            for f in z['families']:
                kids = ', '.join(f"{c['name']}({c['year']},{c['gender']})" for c in f['children'])
                print(f"   {f['role']:<5} {f['husband_name']}({f['husband_birth']}) {f['husband_phone']} / "
                      f"{f['wife_name']}({f['wife_birth']}) {f['wife_phone']} | {kids} | "
                      f"{f['ref_att_m']} {f['ref_att_f']}")
    if not args.dry_run:
        import_zones(zones, reset=args.reset)
        print(f'완료: {DATABASE}')


if __name__ == '__main__':
    main()
