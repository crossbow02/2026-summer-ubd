"""
구역편성 기초자료(가정별 동정) 수집 서비스.

- 구역장: /zone/s/<token>  (카톡으로 받은 링크, 로그인 없음)
- 관리자: /zone/admin      (요약 앱 관리자 로그인 또는 ZONE_ADMIN_PASSWORD)
"""
import io
import json
import os
from datetime import datetime

from flask import (Blueprint, abort, g, jsonify, redirect, render_template, request,
                   send_file, session, url_for)

from zone_survey.db import connect

TARGET_YEAR = int(os.environ.get('ZONE_TARGET_YEAR', datetime.now().year + 1))

ATTENDANCE_LEVELS = [
    ('A', '거의 매번 참석', '구역모임·예배에 꾸준히 참석 (75% 이상)'),
    ('B', '자주 참석', '대체로 참석 (50~75%)'),
    ('C', '가끔 참석', '종종 빠짐 (25~50%)'),
    ('D', '거의 불참', '거의 나오지 못함 (25% 미만)'),
]
NEXT_STATUS = ['구역 배정', '사랑구역', '편성 제외']

bp = Blueprint('zone', __name__, url_prefix='/zone', template_folder='templates')


def db():
    if 'zone_db' not in g:
        g.zone_db = connect()
    return g.zone_db


@bp.teardown_app_request
def _close(exc):
    conn = g.pop('zone_db', None)
    if conn is not None:
        conn.close()


@bp.after_request
def _no_index(resp):
    resp.headers['X-Robots-Tag'] = 'noindex, nofollow'
    resp.headers['Cache-Control'] = 'no-store'
    return resp


def family_to_dict(row):
    d = dict(row)
    d['children'] = json.loads(d.pop('children_json') or '[]')
    d['original'] = json.loads(d.pop('original_json') or '{}')
    d['complete'] = bool(d['info_confirmed'] and d['attendance'])
    d['title'] = ' · '.join(n for n in (d['husband_name'], d['wife_name']) if n) or '(이름 없음)'
    d['changed'] = changed_fields(d)
    return d


def changed_fields(f):
    orig = f['original']
    if not orig:
        return []
    labels = {'husband_name': '형제 이름', 'husband_birth': '형제 생년', 'husband_phone': '형제 연락처',
              'wife_name': '자매 이름', 'wife_birth': '자매 생년', 'wife_phone': '자매 연락처'}
    out = [label for key, label in labels.items() if (f.get(key) or '') != (orig.get(key) or '')]
    if f['children'] != orig.get('children', []):
        out.append('자녀')
    return out


def load_zone(token):
    zone = db().execute('SELECT * FROM zones WHERE token = ?', (token,)).fetchone()
    if not zone:
        abort(404)
    return zone


def load_families(zone_id):
    rows = db().execute('SELECT * FROM zone_families WHERE zone_id = ? ORDER BY sort, id', (zone_id,)).fetchall()
    return [family_to_dict(r) for r in rows]


def clean_children(raw):
    children = []
    for c in raw or []:
        name = str(c.get('name', '')).strip()[:30]
        if not name:
            continue
        gender = c.get('gender') if c.get('gender') in ('남', '여') else ''
        children.append({'name': name, 'year': str(c.get('year', '')).strip()[:10],
                         'gender': gender, 'note': str(c.get('note', '')).strip()[:50]})
    return children


# ---------------------------------------------------------------- 구역장 화면

@bp.route('/s/<token>')
def leader(token):
    zone = load_zone(token)
    families = load_families(zone['id'])
    return render_template('zone/leader.html', zone=zone, families=families, year=TARGET_YEAR,
                           levels=ATTENDANCE_LEVELS, next_status=NEXT_STATUS)


@bp.route('/s/<token>/family/<int:family_id>', methods=['POST'])
def save_family(token, family_id):
    zone = load_zone(token)
    fam = db().execute('SELECT id FROM zone_families WHERE id = ? AND zone_id = ?',
                       (family_id, zone['id'])).fetchone()
    if not fam:
        abort(404)
    data = request.get_json(force=True, silent=True) or {}
    text = lambda k, n=40: str(data.get(k) or '').strip()[:n]
    attendance = data.get('attendance') if data.get('attendance') in ('A', 'B', 'C', 'D') else None
    next_status = data.get('next_status') if data.get('next_status') in NEXT_STATUS else None
    db().execute(
        '''UPDATE zone_families SET husband_name=?, husband_birth=?, husband_phone=?,
           wife_name=?, wife_birth=?, wife_phone=?, children_json=?, info_confirmed=?,
           attendance=?, next_status=?, note=?, updated_at=? WHERE id=?''',
        (text('husband_name'), text('husband_birth', 20), text('husband_phone', 20),
         text('wife_name'), text('wife_birth', 20), text('wife_phone', 20),
         json.dumps(clean_children(data.get('children')), ensure_ascii=False),
         1 if data.get('info_confirmed') else 0, attendance, next_status,
         str(data.get('note') or '').strip()[:3000],
         datetime.now().strftime('%Y-%m-%d %H:%M'), family_id))
    db().commit()
    row = db().execute('SELECT * FROM zone_families WHERE id = ?', (family_id,)).fetchone()
    return jsonify({'ok': True, 'family': family_to_dict(row)})


@bp.route('/s/<token>/family/new', methods=['POST'])
def add_family(token):
    zone = load_zone(token)
    sort = db().execute('SELECT COALESCE(MAX(sort), 0) + 1 FROM zone_families WHERE zone_id = ?',
                        (zone['id'],)).fetchone()[0]
    cur = db().execute(
        "INSERT INTO zone_families (zone_id, sort, role_prev, added_by_leader) VALUES (?, ?, '신규 추가', 1)",
        (zone['id'], sort))
    db().commit()
    return jsonify({'ok': True, 'id': cur.lastrowid})


@bp.route('/s/<token>/family/<int:family_id>/delete', methods=['POST'])
def delete_family(token, family_id):
    zone = load_zone(token)
    # 시트에서 불러온 가정은 지우지 않고 '편성 제외'로 표시하도록 한다
    db().execute('DELETE FROM zone_families WHERE id = ? AND zone_id = ? AND added_by_leader = 1',
                 (family_id, zone['id']))
    db().commit()
    return jsonify({'ok': True})


@bp.route('/s/<token>/submit', methods=['POST'])
def submit_zone(token):
    zone = load_zone(token)
    db().execute('UPDATE zones SET submitted_at = ? WHERE id = ?',
                 (datetime.now().strftime('%Y-%m-%d %H:%M'), zone['id']))
    db().commit()
    return jsonify({'ok': True})


# ---------------------------------------------------------------- 관리자 화면

def is_admin():
    return session.get('role') == 'admin' or session.get('zone_admin')


@bp.route('/admin/login', methods=['GET', 'POST'])
def admin_login():
    error = None
    if request.method == 'POST':
        password = os.environ.get('ZONE_ADMIN_PASSWORD')
        if password and request.form.get('password') == password:
            session['zone_admin'] = True
            return redirect(url_for('zone.admin'))
        error = '비밀번호가 맞지 않습니다.' if password else 'ZONE_ADMIN_PASSWORD 환경변수가 설정되어 있지 않습니다.'
    return render_template('zone/admin_login.html', error=error)


@bp.route('/admin')
def admin():
    if not is_admin():
        return redirect(url_for('zone.admin_login'))
    zones = []
    for z in db().execute('SELECT * FROM zones ORDER BY sort, id').fetchall():
        fams = load_families(z['id'])
        zones.append({**dict(z), 'total': len(fams), 'done': sum(f['complete'] for f in fams),
                      'link': url_for('zone.leader', token=z['token'], _external=True)})
    return render_template('zone/admin.html', zones=zones, year=TARGET_YEAR)


@bp.route('/admin/zone/<int:zone_id>')
def admin_zone(zone_id):
    if not is_admin():
        return redirect(url_for('zone.admin_login'))
    zone = db().execute('SELECT * FROM zones WHERE id = ?', (zone_id,)).fetchone()
    if not zone:
        abort(404)
    return render_template('zone/admin_zone.html', zone=zone, families=load_families(zone_id),
                           year=TARGET_YEAR, link=url_for('zone.leader', token=zone['token'], _external=True))


@bp.route('/admin/export.xlsx')
def admin_export():
    if not is_admin():
        return redirect(url_for('zone.admin_login'))
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = f'{TARGET_YEAR} 구역편성 기초자료'
    headers = ['NO', '현재 구역', '구분', '이름', '생년월일', '연락처', '자녀', '출석률',
               '내년 배정', '동정', '정보 수정', '작성일시', '기존 시트 메모']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='E8EEF7')
    wrap = Alignment(wrap_text=True, vertical='top')
    no = 0
    for z in db().execute('SELECT * FROM zones ORDER BY sort, id').fetchall():
        for f in load_families(z['id']):
            no += 1
            kids = ', '.join(
                f"{c['name']}({','.join(x for x in (c['gender'], c['year']) if x)})" if (c['gender'] or c['year'])
                else c['name'] for c in f['children'])
            people = [(f['husband_name'], f['husband_birth'], f['husband_phone']),
                      (f['wife_name'], f['wife_birth'], f['wife_phone'])]
            people = [p for p in people if p[0]] or [('', '', '')]
            start = ws.max_row + 1
            for i, (name, birth, phone) in enumerate(people):
                family_cols = [no, z['name'], f['role_prev'] or ''] if i == 0 else ['', '', '']
                tail = ([kids, f['attendance'] or '', f['next_status'] or '', f['note'] or '',
                         ', '.join(f['changed']), f['updated_at'] or '', f['sheet_memo'] or '']
                        if i == 0 else [''] * 7)
                ws.append(family_cols + [name, birth, phone] + tail)
            end = ws.max_row
            if end > start:
                for col in (1, 2, 3, 7, 8, 9, 10, 11, 12, 13):
                    ws.merge_cells(start_row=start, start_column=col, end_row=end, end_column=col)
            for row in ws.iter_rows(min_row=start, max_row=end):
                for cell in row:
                    cell.alignment = wrap
    for col, width in zip('ABCDEFGHIJKLM', (5, 9, 9, 10, 11, 15, 32, 7, 10, 60, 14, 16, 40)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f'{TARGET_YEAR}_구역편성_기초자료_{datetime.now():%m%d}.xlsx'
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
